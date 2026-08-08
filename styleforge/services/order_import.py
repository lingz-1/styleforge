"""Parse shopping exports into privacy-safe wardrobe import candidates."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import BinaryIO
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit


HEADER_ALIASES = {
    "external_order_id": ("订单号", "订单编号", "交易号"),
    "order_submitted_at": ("订单提交时间", "下单时间", "订单时间", "创建时间"),
    "order_status": ("订单状态", "交易状态"),
    "shop_name": ("店铺名称", "商家名称", "卖家名称", "店铺"),
    "refund_status": ("退款状态", "退款/售后状态", "退款售后状态"),
    "after_sale_status": ("售后状态", "退货状态", "维权状态"),
    "logistics_status": ("物流状态", "配送状态", "收货状态"),
    "product_name": ("商品名称", "商品标题", "宝贝名称", "商品", "标题"),
    "product_url": ("商品链接", "宝贝链接", "链接", "商品网址"),
    "variant_text": ("型号款式", "商品规格", "规格", "款式", "SKU"),
    "quantity": ("商品数量", "数量", "购买数量"),
    "listed_amount": ("商品金额", "商品价格", "单价", "金额"),
    "paid_amount": ("实付金额", "实付款", "付款金额", "实际支付"),
}

PARSER_REVISION = "wardrobe-orders-v3"

RECEIVED_ORDER_STATUSES = {
    "交易成功",
    "已完成",
    "订单完成",
    "已收货",
    "确认收货",
    "已签收",
}
NEUTRAL_AFTER_SALE_STATUSES = {
    "",
    "无退款",
    "未退款",
    "无售后",
    "未申请售后",
    "退款关闭",
    "售后关闭",
}
DISALLOWED_STATUS_PATTERN = re.compile(
    r"退款成功|退款中|申请退款|退货|交易关闭|订单关闭|已取消|取消订单|"
    r"售后中|维权中|拒收|待发货|已付款|卖家已发货|运输中|派送中|待收货",
    re.IGNORECASE,
)

NON_FASHION_PATTERN = re.compile(
    r"牙膏|饮用水|矿泉水|护肤|水乳|面膜|洗发|沐浴|键盘|键帽|鼠标|"
    r"文件夹|卡册|内页|收纳册|扇子|摆件|手机壳|数据线|充电|零食|食品|"
    r"清洁剂|卫生纸|纸巾|打印|胶带|药品|保健品|化妆品|口红|粉底|香水|"
    r"卫生巾|安睡裤|夜安裤|姨妈巾|消毒液|除菌液|消毒水|杀菌液|"
    r"娃衣|棉花娃娃|玩偶|公仔|钥匙扣|挂件|"
    r"主板|cpu|显卡|内存条|处理器|intel|amd|"
    r"修眉|眉刀|美甲|指甲|搓条|粉扑|"
    r"睡袋|床单|被套|枕套|拼豆|材料包|"
    r"刻刀|雕刻|剪纸|美工刀|刀片|裁纸刀|剪刀|"
    r"首饰盒|戒指盒|珠宝盒|收纳盒|道具|cosplay",
    re.IGNORECASE,
)


@dataclass(frozen=True, slots=True)
class CategoryRule:
    item_type: str
    subtype: str
    pattern: re.Pattern[str]


CATEGORY_RULES = (
    CategoryRule("underwear", "bra", re.compile(r"文胸|胸罩|内衣|bralette", re.I)),
    CategoryRule(
        "underwear", "underpants", re.compile(r"内裤|三角裤|平角裤|安全裤", re.I)
    ),
    CategoryRule("sleepwear", "sleepwear", re.compile(r"睡衣|睡裙|睡袍|家居服", re.I)),
    CategoryRule("shoes", "boots", re.compile(r"靴|boot", re.I)),
    CategoryRule("shoes", "loafers", re.compile(r"乐福鞋|loafer", re.I)),
    CategoryRule("shoes", "flats", re.compile(r"玛丽珍|平底鞋|单鞋", re.I)),
    CategoryRule("shoes", "sneakers", re.compile(r"运动鞋|跑步鞋|老爹鞋|板鞋|帆布鞋", re.I)),
    CategoryRule("shoes", "sandals", re.compile(r"凉鞋|拖鞋|凉拖", re.I)),
    CategoryRule("shoes", "shoes", re.compile(r"女鞋|男鞋|皮鞋|鞋子|鞋", re.I)),
    CategoryRule("bag", "bag", re.compile(r"单肩包|双肩包|手提包|托特包|腋下包|挎包|背包|钱包|卡包", re.I)),
    CategoryRule("eyewear", "eyewear", re.compile(r"太阳镜|墨镜|眼镜框|眼镜", re.I)),
    CategoryRule("earrings", "earrings", re.compile(r"耳环|耳钉|耳夹|耳饰", re.I)),
    CategoryRule("necklace", "necklace", re.compile(r"项链|颈链", re.I)),
    CategoryRule("bracelet", "bracelet", re.compile(r"手链|手镯", re.I)),
    CategoryRule("rings", "rings", re.compile(r"戒指|指环", re.I)),
    CategoryRule("belts", "belts", re.compile(r"腰带|皮带", re.I)),
    CategoryRule("hats", "hats", re.compile(r"帽子|鸭舌帽|渔夫帽|贝雷帽|针织帽", re.I)),
    CategoryRule("hairwear", "hairwear", re.compile(r"发夹|发箍|发带|头饰", re.I)),
    CategoryRule("jewellery", "jewellery", re.compile(r"饰品|配饰|胸针", re.I)),
    CategoryRule("dress", "qipao", re.compile(r"旗袍", re.I)),
    CategoryRule("dress", "midi_dress", re.compile(r"中长连衣裙|中长裙", re.I)),
    CategoryRule("dress", "dress", re.compile(r"连衣裙|礼服裙|礼服", re.I)),
    CategoryRule("jumpsuit", "jumpsuit", re.compile(r"连体裤|连身裤|jumpsuit", re.I)),
    CategoryRule("suit", "suit", re.compile(r"西装套装|正装套装|礼服套装", re.I)),
    CategoryRule("outfit_set", "outfit_set", re.compile(r"两件套|三件套|四件套|全套|套装", re.I)),
    CategoryRule("outwear", "coat", re.compile(r"大衣|呢子外套", re.I)),
    CategoryRule("outwear", "down_jacket", re.compile(r"羽绒服|棉服", re.I)),
    CategoryRule("outwear", "blazer", re.compile(r"西装外套|西服外套|西服", re.I)),
    CategoryRule("outwear", "jacket", re.compile(r"夹克|冲锋衣|风衣|外套", re.I)),
    CategoryRule("outwear", "cardigan", re.compile(r"开衫|褙子|长衫|披肩", re.I)),
    CategoryRule("skirt", "pleated_skirt", re.compile(r"百褶裙", re.I)),
    CategoryRule("skirt", "midi_skirt", re.compile(r"中长半身裙|中长裙", re.I)),
    CategoryRule("skirt", "mini_skirt", re.compile(r"短裙|迷你裙", re.I)),
    CategoryRule("skirt", "skirt", re.compile(r"半身裙|下裙|马面裙|长裙|裙子", re.I)),
    CategoryRule("shorts", "shorts", re.compile(r"短裤", re.I)),
    CategoryRule("pants", "jeans", re.compile(r"牛仔裤", re.I)),
    CategoryRule("pants", "leggings", re.compile(r"瑜伽裤|紧身裤|打底裤", re.I)),
    CategoryRule("pants", "tailored_trousers", re.compile(r"西裤|正装裤", re.I)),
    CategoryRule(
        "pants",
        "pants",
        re.compile(r"中裤|七分裤|九分裤|长裤|阔腿裤|直筒裤|休闲裤|裤子|裤", re.I),
    ),
    CategoryRule("top", "shirt", re.compile(r"衬衫|衬衣|长衫", re.I)),
    CategoryRule("top", "tank_top", re.compile(r"背心|吊带|小背心|无袖上衣", re.I)),
    CategoryRule("top", "t_shirt", re.compile(r"T恤|t恤|tee", re.I)),
    CategoryRule("top", "knitwear", re.compile(r"毛衣|针织衫|针织上衣", re.I)),
    CategoryRule("top", "hoodie", re.compile(r"卫衣|连帽衫", re.I)),
    CategoryRule("top", "top", re.compile(r"上衣|短袖|长袖|罩衫|Polo衫|polo", re.I)),
    CategoryRule("legwear", "socks", re.compile(r"袜子|长筒袜|丝袜|压力袜", re.I)),
)

COLOR_ALIASES = (
    "黑色",
    "白色",
    "灰色",
    "银色",
    "金色",
    "红色",
    "酒红",
    "粉色",
    "粉",
    "蓝色",
    "浅蓝",
    "深蓝",
    "湖蓝",
    "藏青",
    "绿色",
    "浅绿",
    "薄荷绿",
    "黄色",
    "紫色",
    "棕色",
    "咖色",
    "卡其",
    "米白",
    "杏色",
    "奶杏",
)


@dataclass(frozen=True, slots=True)
class ParsedOrderRow:
    source_row_number: int
    row_hash: str
    platform: str
    external_order_id_hash: str
    order_submitted_at: str
    order_status: str
    shop_name: str
    refund_status: str
    after_sale_status: str
    logistics_status: str
    order_eligibility: str
    order_eligibility_reason: str
    external_product_id: str
    product_name: str
    canonical_url: str
    variant_text: str
    quantity: int
    listed_amount: float | None
    paid_amount: float | None
    currency: str
    predicted_item_type: str
    predicted_subtype: str
    predicted_color: str
    predicted_size: str
    predicted_audience: str
    confidence: float
    decision: str
    decision_reason: str

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True, slots=True)
class ParsedOrderWorkbook:
    file_sha256: str
    platform: str
    sheet_name: str
    rows: tuple[ParsedOrderRow, ...]

    def statistics(self) -> dict[str, object]:
        candidate_count = sum(row.decision == "candidate" for row in self.rows)
        return {
            "total_rows": len(self.rows),
            "candidate_rows": candidate_count,
            "excluded_rows": len(self.rows) - candidate_count,
            "rows_with_paid_amount": sum(row.paid_amount is not None for row in self.rows),
            "rows_with_valid_product_id": sum(bool(row.external_product_id) for row in self.rows),
            "rows_with_order_id": sum(bool(row.external_order_id_hash) for row in self.rows),
            "eligible_order_rows": sum(
                row.order_eligibility == "eligible" for row in self.rows
            ),
            "ineligible_order_rows": sum(
                row.order_eligibility == "ineligible" for row in self.rows
            ),
            "unknown_order_rows": sum(
                row.order_eligibility == "unknown" for row in self.rows
            ),
            "by_order_status": _counts(row.order_status or "unknown" for row in self.rows),
            "by_order_eligibility": _counts(row.order_eligibility for row in self.rows),
            "by_item_type": _counts(
                row.predicted_item_type or "unclassified" for row in self.rows
            ),
        }


def _counts(values) -> dict[str, int]:
    result: dict[str, int] = {}
    for value in values:
        result[value] = result.get(value, 0) + 1
    return dict(sorted(result.items(), key=lambda item: (-item[1], item[0])))


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalized_header(value: object) -> str:
    return re.sub(r"\s+", "", _text(value)).lower()


def _parse_money(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    normalized = re.sub(r"[￥¥,\s]", "", _text(value))
    if not normalized:
        return None
    try:
        return float(normalized)
    except ValueError:
        return None


def _parse_quantity(value: object) -> int:
    try:
        parsed = int(float(_text(value)))
    except ValueError:
        return 1
    return max(parsed, 1)


def _canonicalize_url(value: object) -> tuple[str, str, str]:
    raw_url = _text(value)
    try:
        parsed = urlsplit(raw_url)
    except ValueError:
        return "unknown", "", ""
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return "unknown", "", ""
    host = parsed.hostname.lower()
    if "taobao.com" in host or "tmall.com" in host:
        platform = "taobao"
        identifier = parse_qs(parsed.query).get("id", [""])[0]
        query = urlencode({"id": identifier}) if identifier else ""
        return platform, identifier, urlunsplit(("https", host, parsed.path, query, ""))
    if "jd.com" in host:
        match = re.search(r"/(\d+)\.html", parsed.path)
        identifier = match.group(1) if match else ""
        return "jd", identifier, urlunsplit(("https", host, parsed.path, "", ""))
    if "pinduoduo.com" in host:
        identifier = parse_qs(parsed.query).get("goods_id", [""])[0]
        query = urlencode({"goods_id": identifier}) if identifier else ""
        return "pinduoduo", identifier, urlunsplit(("https", host, parsed.path, query, ""))
    return host, "", urlunsplit(("https", host, parsed.path, "", ""))


def _match_category(product_name: str, variant_text: str) -> tuple[str, str, float, str]:
    if NON_FASHION_PATTERN.search(f"{product_name} {variant_text}"):
        return "", "", 0.95, "explicit_non_fashion_signal"
    for source_name, source_text, confidence in (
        ("variant", variant_text, 0.9),
        ("title", product_name, 0.76),
    ):
        for rule in CATEGORY_RULES:
            if rule.pattern.search(source_text):
                return rule.item_type, rule.subtype, confidence, f"matched_{source_name}"
    return "", "", 0.0, "no_wearable_signal"


def _extract_color(product_name: str, variant_text: str) -> str:
    combined = f"{variant_text} {product_name}"
    for color in COLOR_ALIASES:
        if color.lower() in combined.lower():
            return color
    return ""


def _extract_size(variant_text: str) -> str:
    patterns = (
        r"(?:^|[;；,，\s])((?:XXS|XS|S|M|L|XL|XXL|XXXL|\d{2,3}/\d{2,3}[A-Z]?))(?:$|[;；,，\s\[(])",
        r"(?:^|[;；,，\s])(均码|F码|F)(?:$|[;；,，\s\[(])",
        r"(?:^|[;；,，\s])(3[4-9]|4[0-8])(?:$|[;；,，\s\[(])",
    )
    for pattern in patterns:
        match = re.search(pattern, variant_text, re.I)
        if match:
            return match.group(1).upper()
    return ""


def _infer_audience(product_name: str, default_audience: str) -> str:
    if re.search(r"女童|女孩|儿童女", product_name, re.I):
        return "girls"
    if re.search(r"男童|男孩|儿童男", product_name, re.I):
        return "boys"
    if re.search(r"婴儿|宝宝|婴童", product_name, re.I):
        return "baby"
    if re.search(r"男士|男款|男装|男人|男子", product_name, re.I):
        return "men"
    if re.search(r"女士|女款|女装|女人|女子", product_name, re.I):
        return "women"
    return default_audience


def _row_hash(values: tuple[object, ...]) -> str:
    payload = "\x1f".join(_text(value) for value in values)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _order_id_hash(order_id: str) -> str:
    if not order_id:
        return ""
    return hashlib.sha256(order_id.encode("utf-8")).hexdigest()


def _evaluate_order_eligibility(
    *,
    order_status: str,
    refund_status: str,
    after_sale_status: str,
    logistics_status: str,
) -> tuple[str, str]:
    for field_name, status in (
        ("refund_status", refund_status),
        ("after_sale_status", after_sale_status),
    ):
        if status in NEUTRAL_AFTER_SALE_STATUSES:
            continue
        if DISALLOWED_STATUS_PATTERN.search(status) or status:
            return "ineligible", f"{field_name}:{status}"
    if DISALLOWED_STATUS_PATTERN.search(logistics_status):
        return "ineligible", f"logistics_status:{logistics_status}"
    if order_status in RECEIVED_ORDER_STATUSES:
        return "eligible", "received_without_refund_signal"
    if order_status:
        return "ineligible", f"order_status_not_received:{order_status}"
    return "unknown", "missing_order_status"


def _read_bytes(source: Path | bytes | BinaryIO) -> bytes:
    if isinstance(source, Path):
        return source.read_bytes()
    if isinstance(source, bytes):
        return source
    return source.read()


def parse_order_workbook(
    source: Path | bytes | BinaryIO,
    *,
    default_audience: str = "",
) -> ParsedOrderWorkbook:
    """Read an order workbook without persisting unapproved raw columns."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel import requires openpyxl. Install the StyleForge orders extra."
        ) from error

    if default_audience not in {"", "women", "men", "girls", "boys", "baby", "life"}:
        raise ValueError(f"Unsupported default audience: {default_audience}")
    data = _read_bytes(source)
    if not data:
        raise ValueError("Order workbook is empty")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        selected = None
        column_map: dict[str, int] = {}
        header_row_number = 0
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1
            ):
                normalized = {
                    _normalized_header(value): index
                    for index, value in enumerate(row)
                    if _normalized_header(value)
                }
                current_map = {}
                for field, aliases in HEADER_ALIASES.items():
                    for alias in aliases:
                        key = _normalized_header(alias)
                        if key in normalized:
                            current_map[field] = normalized[key]
                            break
                if "product_name" in current_map:
                    selected = sheet
                    column_map = current_map
                    header_row_number = row_number
                    break
            if selected is not None:
                break
        if selected is None:
            raise ValueError("No worksheet contains a recognized product-name column")

        parsed_rows = []
        platform_counts: dict[str, int] = {}
        order_context = {
            "external_order_id": "",
            "order_submitted_at": "",
            "order_status": "",
            "shop_name": "",
            "refund_status": "",
            "after_sale_status": "",
            "logistics_status": "",
        }
        for row_number, values in enumerate(
            selected.iter_rows(min_row=header_row_number + 1, values_only=True),
            start=header_row_number + 1,
        ):
            def value(field: str) -> object:
                index = column_map.get(field)
                return values[index] if index is not None and index < len(values) else None

            product_name = _text(value("product_name"))
            if not product_name:
                continue
            current_order_values = {
                field: _text(value(field))
                for field in order_context
            }
            if current_order_values["external_order_id"]:
                order_context = current_order_values
            else:
                for field, current_value in current_order_values.items():
                    if current_value:
                        order_context[field] = current_value
            variant_text = _text(value("variant_text"))
            platform, product_id, canonical_url = _canonicalize_url(value("product_url"))
            platform_counts[platform] = platform_counts.get(platform, 0) + 1
            item_type, subtype, confidence, reason = _match_category(
                product_name, variant_text
            )
            eligibility, eligibility_reason = _evaluate_order_eligibility(
                order_status=order_context["order_status"],
                refund_status=order_context["refund_status"],
                after_sale_status=order_context["after_sale_status"],
                logistics_status=order_context["logistics_status"],
            )
            if eligibility != "eligible":
                decision = "excluded"
                decision_reason = eligibility_reason
            elif item_type:
                decision = "candidate"
                decision_reason = reason
            else:
                decision = "excluded"
                decision_reason = reason
            parsed_rows.append(
                ParsedOrderRow(
                    source_row_number=row_number,
                    row_hash=_row_hash(tuple(values)),
                    platform=platform,
                    external_order_id_hash=_order_id_hash(
                        order_context["external_order_id"]
                    ),
                    order_submitted_at=order_context["order_submitted_at"],
                    order_status=order_context["order_status"],
                    shop_name=order_context["shop_name"],
                    refund_status=order_context["refund_status"],
                    after_sale_status=order_context["after_sale_status"],
                    logistics_status=order_context["logistics_status"],
                    order_eligibility=eligibility,
                    order_eligibility_reason=eligibility_reason,
                    external_product_id=product_id,
                    product_name=product_name,
                    canonical_url=canonical_url,
                    variant_text=variant_text,
                    quantity=_parse_quantity(value("quantity")),
                    listed_amount=_parse_money(value("listed_amount")),
                    paid_amount=_parse_money(value("paid_amount")),
                    currency="CNY",
                    predicted_item_type=item_type,
                    predicted_subtype=subtype,
                    predicted_color=_extract_color(product_name, variant_text),
                    predicted_size=_extract_size(variant_text),
                    predicted_audience=_infer_audience(product_name, default_audience),
                    confidence=confidence,
                    decision=decision,
                    decision_reason=decision_reason,
                )
            )
        if not parsed_rows:
            raise ValueError("Order workbook contains no product rows")
        platform = max(platform_counts, key=platform_counts.get)
        return ParsedOrderWorkbook(
            file_sha256=hashlib.sha256(data).hexdigest(),
            platform=platform,
            sheet_name=selected.title,
            rows=tuple(parsed_rows),
        )
    finally:
        workbook.close()
