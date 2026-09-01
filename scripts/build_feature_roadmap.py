from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


STATUS_ROWS = [
    ("核心 Agent", "多 Agent 推荐主流程", "已完成", 100, "StyleForgeHarness；Coordinator→Research→Stylist×3→Critic→持久化；完整质量门禁通过", "保持现架构，不再增加同职责 Agent", "继续通过 typed state、工具权限和评测增强", "维持"),
    ("核心 Agent", "五维评价与结果合同", "已完成", 100, "五维 ScoreCard、统一结果合同、本地/Agentic 双路径归一；729 项测试通过", "保持并继续扩展固定 benchmark", "把五维阈值覆盖扩展到更多真实请求", "P0"),
    ("核心 Agent", "六类任务路由与扩展业务", "已完成", 100, "推荐/修改/风格建议/单品建议/兼容性/衣柜缺口，共享 POST /tasks/execute", "保持", "在统一 benchmark 中按任务分别统计", "P0"),
    ("核心 Agent", "语义/混合/关键词降级检索", "已完成", 100, "FashionCLIP 512 维 CUDA 实测；semantic、hybrid、keyword fallback 均通过", "保持并量化效果", "增加同一真值集上的召回率、槽位命中率对比", "P0"),
    ("核心 Agent", "推荐质量固定基准", "已完成基础闭环", 90, "已建 7 类固定衣橱、60 条分层案例和 2 条 Polyvore 真实案例；42 条轻量子集 42/42；报告含按任务/衣橱分布和失败明细", "开始分批建立 18 条 configured 六任务真实模型基线", "补多轮、天气、权重和失败恢复案例；真实 LLM/GPU 保持显式运行", "P0"),
    ("核心 Agent", "EXT-002：one_piece 完整性", "已完成", 100, "显式鞋/外套/配饰进入检索和硬校验；Polyvore 真实婚礼/约会 + DeepSeek 2/2 通过；单请求模型调用降为 1 次", "保持真实基准回归", "继续扩大场景覆盖，保留 grounded fallback 作为安全兜底", "维持"),
    ("核心 Agent", "EXT-003：记忆泛化边界", "已完成", 100, "普通替换不归纳品类；同场景 3 个不同单品且 3 次不同判断才生成弱 category 假设；跨场景与重复事件隔离", "保持并加入长期回归", "继续用真实多轮案例验证显式偏好可覆盖弱归纳", "维持"),
    ("衣柜", "上传、图片识别、批量识别", "已完成", 100, "单图 analyze/photo、异步 batch-recognize、进度查询/撤销均有 API 与测试", "保持", "后续补生产任务队列，不改业务合同", "维持"),
    ("衣柜", "订单导入、审核、提交", "已完成（真实个人提交待人工确认）", 90, "预览/幂等/状态门禁/提交完整；真实订单候选仍需用户人工抽查", "需要人工验收，不继续自动提交", "抽查后只提交明确仍持有的衣物", "人工"),
    ("衣柜", "单品 CRUD、分类目录、图片展示", "已完成", 100, "wardrobes、catalog/search、catalog/taxonomy、items/images API 与 Web 页面", "保持", "补 UI 回归即可", "维持"),
    ("衣柜", "衣柜筛选与排序", "部分完成", 45, "已有按品类折叠、搜索候选和基础展示；缺季节/颜色/多条件筛选与用户排序", "质量基准后做的低风险 UX", "前端 query state + 后端可选过滤参数；先本地过滤再按规模下沉", "P1"),
    ("衣柜", "上传搭配图搜索相似穿搭", "未完成", 10, "已有 FashionCLIP 单品向量检索，但没有 outfit 图→衣柜组合检索入口", "有真实需求再做", "图像 embedding→Top-K 单品→按槽位重组→Agent 审校", "P3"),
    ("会话与记忆", "多轮对话与局部修改", "已完成", 100, "chat_sessions/chat_messages、历史恢复、整体调整、槽位替换均有实现与测试", "继续补反悔还原", "保存每轮 outfit diff，支持“换回来/其余保持”", "P1"),
    ("会话与记忆", "反馈闭环与长期偏好", "已完成", 100, "点赞/踩/采纳/换掉→evidence→preference_model→后续 Agent 注入；Web 可编辑/遗忘", "保持，重点修泛化边界", "用固定跨会话用例验证偏好生效与用户隔离", "P0"),
    ("会话与记忆", "收藏/喜爱搭配", "已完成", 90, "保存、列表、删除 API 与前端“添加至穿搭集”；支持将保存动作写入记忆", "补独立管理页可提升完整度", "复用 saved-outfits API 增加查看/删除/基于此修改", "P1"),
    ("会话与记忆", "会话统计/数据看板", "未完成", 10, "已有事件与运行数据，但没有用户统计页面", "低优先级", "先定义真正有用的 4–6 个指标，避免做装饰性看板", "P3"),
    ("天气与时空", "地点回退与设备定位", "已完成", 95, "显式地点→设备坐标→默认地点；坐标精度、时效、授权失败均有测试", "保持", "补用户级默认城市持久化即可", "P1"),
    ("天气与时空", "实时天气与小时事实", "已完成基础闭环", 85, "Open-Meteo、daily/hourly、体感/降水/风/UV 等 schema 与 Web WeatherCard", "纳入场景 benchmark", "对明天、晚场、雨天、查不到四类做事实引用断言", "P1"),
    ("天气与时空", "复杂时间表达与跨天窗口", "部分完成", 55, "typed requirements 已支持窗口/时段，但节日、复杂相对日期和跨城市行程仍不完整", "在推荐质量稳定后分期做", "独立 TimeWindow parser；返回 granularity/precision/approximate/source", "P2"),
    ("天气与时空", "Event Lookup/场馆天气", "未完成", 15, "已有 web_search、event grounding 与技能流程，但无稳定事件/场次解析器", "不作为当前下一步", "事件搜索→证据抽取→场馆/时间确认→小时天气；不确定时澄清", "P3"),
    ("天气与时空", "历史气候参考", "未完成", 0, "远期请求目前不能提供可验证多年统计", "低于 Event Lookup", "接历史气候数据并显式标注 climate_reference，禁止伪装预报", "P4"),
    ("知识与搜索", "知识 RAG 向量化", "已完成基础闭环", 85, "Markdown 分块、FashionCLIP 文本 embedding、ChromaStore、混合召回与关键词降级已有实现", "原表描述已过期", "补离线索引健康检查和知识命中评测", "P1"),
    ("知识与搜索", "通用联网搜索", "已完成（需 Tavily 配置）", 80, "Agentic web_search 工具、grounding search-before-ask 已接入", "保持工具边界", "只为待查事实调用，保留来源和失败语义", "维持"),
    ("知识与搜索", "电商商品搜索与链接", "未完成", 15, "衣柜缺口能指出缺什么，但没有稳定商品结果、价格和链接合同", "明确需要购买建议时再做", "缺口结构→关键词→电商搜索→去重/价格过滤→来源链接；与衣柜推荐分离", "P3"),
    ("个性化", "五维权重与设置", "已完成", 100, "用户 evaluation preference API 与 Web 设置页", "保持", "纳入 benchmark 检查排序确实随权重变化", "P0"),
    ("个性化", "身体/面部/肤色画像", "未完成", 0, "无实现；涉及敏感生物与身体数据", "目前没必要做", "如未来做，只允许用户主动填写低敏属性，避免自动面部推断", "不做"),
    ("造型输出", "发型、妆容、随身物品建议", "部分完成", 70, "结果合同支持 carry/advice，Agent 可生成；尚缺独立、稳定的结构化全造型合同", "先通过 benchmark 验证价值", "新增 grooming/accessory 建议字段并由 Critic 校验事实依据", "P2"),
    ("造型输出", "虚拟试穿", "未完成", 0, "无模型、素材授权、显存和延迟方案", "当前没必要做", "应作为独立项目评估，不嵌入现有推荐主链路", "不做"),
    ("工程化", "完整本地质量门禁", "已完成", 100, "729 测试+42 条轻量推荐质量基准+构建+浏览器 E2E+真实 CUDA+独立 PostgreSQL schema 清理", "保持", "真实 LLM/GPU 评测继续显式按需运行", "维持"),
    ("工程化", "CI", "已完成基础闭环", 90, "GitHub Actions 已接入 PostgreSQL、Ruff、无 GPU 全量 pytest、42 条轻量质量基准、API health smoke 和 Vite build", "观察首轮远程运行并按耗时优化", "真实 LLM/GPU 保持显式本地门禁，不在普通 PR 消耗", "P1"),
    ("工程化", "Windows 一键启动/Docker", "已完成启动器", 75, "PowerShell 统一支持 Check/Start/Status/Stop，检查 Python/依赖/PostgreSQL并管理 API+Web/Streamlit 日志", "先用真实本机启动验收；Docker 暂缓", "如需部署再封装 API/PG/Redis，数据集继续外部只读挂载", "P1"),
    ("工程化", "WebSocket 实时推送", "未完成", 0, "批量识别当前通过轮询稳定工作", "目前没必要做", "只有当长任务并发和状态延迟成为问题时再引入 SSE/WebSocket", "不做"),
    ("工程化", "Prompt 管理界面", "未完成", 20, "Prompt 已文件化并带版本；无可视化后台", "目前不优先", "先建立 prompt/eval 版本绑定；出现非研发运营角色后再做 UI", "P4"),
    ("数据", "Mytheresa 全量主库接入", "代码完成、生产验收未完成", 65, "解析/多图/幂等导入已实现；62,457 商品和 328,754 图片尚未正式验收", "独立于当前推荐质量阶段", "备份→像素审计→一次性库故障演练→全量导入→索引", "P3"),
    ("数据", "Polyvore 兼容性/FITB 离线基线", "部分完成", 55, "真实商品映射、图片核验和 2 条 Agent 验收已闭环；仍缺随机/共现/FashionCLIP 的统一 FITB 报告", "作为推荐质量第二批", "统一 item 映射和指标，产出可复现 polyvore_baselines.json", "P1"),
]


ROADMAP_ROWS = [
    (1, "P0-A", "现状基线冻结", "本工作簿完成；同步文档中的过期状态；整理当前未提交变更", "功能状态有证据、质量门禁报告可追溯、无误称", "本轮"),
    (2, "P0-B", "Wardrobe Fixtures", "已完成六类快速回归衣橱和一类 Polyvore 真实衣橱，使用稳定 UUID", "fixtures 可重复导入；用户隔离；真实商品保留 source/dataset_item_id/图片路径", "已完成"),
    (3, "P0-C", "统一五维 benchmark", "已完成 60 条分层案例；42 条轻量子集 42/42，18 条 configured 覆盖六任务待逐批建立真实模型基线", "报告含逐任务/衣橱成功率、硬约束、五维字段和失败明细；本地可复现", "基础闭环完成"),
    (4, "P0-D", "修复 EXT-002", "显式支撑槽位检索、状态一致性和 grounded fallback 已实现", "真实 Polyvore + DeepSeek 2/2 通过；不再漏鞋/配饰/外套", "已完成"),
    (5, "P0-E", "修复 EXT-003", "普通替换排除、同场景/不同单品/不同事件三重阈值与里程碑去重已实现", "单品负反馈不污染整类；多条一致证据才升级；跨用户严格隔离", "已完成"),
    (6, "P1", "轻量评测接入质量门禁", "42 条无 LLM、无 GPU 子集已接入本地质量门禁和 CI；真实 LLM/GPU 保留显式任务", "日常回归速度可控，关键质量退化能阻断", "已完成"),
    (7, "P2", "CI 与 Windows 一键启动", "GitHub Actions 与 PowerShell launcher 已实现；等待远程首跑和本机启动验收", "PR 自动跑静态检查、测试、构建和 smoke；本地 Check/Start/Status/Stop 可用", "进行中"),
    (8, "P3", "产品增强按需求选择", "衣柜筛选/收藏管理优先；商品搜索、Event Lookup、相似穿搭随后", "每项必须先有用户场景和验收用例，不以功能数量为目标", "后续"),
]


DEFER_ROWS = [
    ("新增更多主 Agent", "当前职责已由 Coordinator/Research/Stylist/Critic 清晰覆盖；增加 Agent 会增加调用成本和状态复杂度", "通过工具、typed state、评测和子图扩展现有职责"),
    ("WebSocket", "批量识别轮询已经稳定，当前没有高并发实时性证据", "达到轮询瓶颈后优先考虑 SSE，再评估 WebSocket"),
    ("虚拟试穿", "模型、隐私、显存、延迟和素材授权均是独立问题，会稀释推荐系统主线", "作为独立实验项目验证，不接主生产链路"),
    ("自动面部/肤色/身体推断", "涉及敏感数据和偏见风险，且不是推荐质量当前瓶颈", "只接受用户主动填写的低敏偏好或尺码信息"),
    ("Prompt 管理后台", "Prompt 已文件化、可版本控制；当前没有非研发运营角色", "先绑定 prompt 版本与 eval 报告，需求成熟后再做 UI"),
    ("完整历史气候系统", "使用频率与价值尚未验证，优先级低于近期天气和事件查询", "先诚实返回超窗与不确定性，后续接可追溯气候数据"),
]


def style_sheet(sheet, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    body_font = Font(name="Microsoft YaHei", size=10)
    thin = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False


def add_status_sheet(workbook) -> None:
    sheet = workbook.create_sheet("功能现状")
    sheet.append(["类别", "功能", "当前状态", "完成度", "证据与边界", "下一步判断", "大致实现", "优先级"])
    for row in STATUS_ROWS:
        sheet.append(row)
    for cell in sheet["D"][1:]:
        cell.number_format = '0"%"'
    status_colors = {
        "已完成": "E2F0D9",
        "部分完成": "FFF2CC",
        "未完成": "FCE4D6",
        "代码完成、生产验收未完成": "FFF2CC",
        "数据与计划具备、执行未闭环": "FFF2CC",
        "已完成基础闭环": "E2F0D9",
        "已完成（需 Tavily 配置）": "E2F0D9",
        "已完成（真实个人提交待人工确认）": "FFF2CC",
    }
    for row in range(2, sheet.max_row + 1):
        value = str(sheet.cell(row, 3).value)
        color = status_colors.get(value)
        if color:
            sheet.cell(row, 3).fill = PatternFill("solid", fgColor=color)
    style_sheet(sheet, [16, 30, 28, 10, 58, 34, 58, 12])


def add_roadmap_sheet(workbook) -> None:
    sheet = workbook.create_sheet("下一步路线图")
    sheet.append(["顺序", "阶段", "任务", "实施内容", "完成判据", "建议时点"])
    for row in ROADMAP_ROWS:
        sheet.append(row)
    style_sheet(sheet, [8, 12, 28, 68, 62, 16])


def add_defer_sheet(workbook) -> None:
    sheet = workbook.create_sheet("暂不优先")
    sheet.append(["功能", "暂不优先原因", "替代方案/重新评估条件"])
    for row in DEFER_ROWS:
        sheet.append(row)
    style_sheet(sheet, [30, 70, 65])


def build(source: Path, output: Path) -> None:
    workbook = load_workbook(source)
    workbook.active.title = "原始预想"
    for name in ("功能现状", "下一步路线图", "暂不优先"):
        if name in workbook.sheetnames:
            del workbook[name]
    add_status_sheet(workbook)
    add_roadmap_sheet(workbook)
    add_defer_sheet(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the current StyleForge feature roadmap workbook.")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
